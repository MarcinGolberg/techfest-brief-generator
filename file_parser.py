import json
import os
from datetime import datetime
from pypdf import PdfReader
import docx
import pptx


class BriefInputParser:
    def __init__(self):
        pass

    def _extract_pdf(self, file_path):
        text = ""
        try:
            reader = PdfReader(file_path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        except Exception as e:
            return f"[BŁĄD ODCZYTU PDF: {e}]"
        return text.strip()

    def _extract_docx(self, file_path):
        text = ""
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text + "\n"
        except Exception as e:
            return f"[BŁĄD ODCZYTU DOCX: {e}]"
        return text.strip()

    def _extract_pptx(self, file_path):
        text = ""
        try:
            prs = pptx.Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text += shape.text + "\n"
        except Exception as e:
            return f"[BŁĄD ODCZYTU PPTX: {e}]"
        return text.strip()

    def generate_json_payload(self, raw_texts=None, file_paths=None, save_to_file=True):
        """
        Konwertuje teksty i pliki do ustrukturyzowanego formatu JSON.
        Zwraca ciąg znaków JSON oraz opcjonalnie zapisuje go do pliku.
        """
        if raw_texts is None:
            raw_texts = []
        if file_paths is None:
            file_paths = []

        sources = []
        combined_text_list = []

        # 1. Przetwarzanie tekstów wklejonych ręcznie / maili
        for text in raw_texts:
            # Prosta heurystyka dla maili
            if any(k in text for k in ["Od:", "From:", "Temat:"]):
                source_type = "email"
            else:
                source_type = "manual_text"

            sources.append({
                "type": source_type,
                "filename": None,
                "content": text
            })
            combined_text_list.append(f"--- Źródło: {source_type.upper()} ---\n{text}")

        # 2. Przetwarzanie plików
        for file_path in file_paths:
            if not os.path.exists(file_path):
                print(f"Ostrzeżenie: Plik '{file_path}' nie istnieje i zostanie pominięty.")
                continue

            filename = os.path.basename(file_path)
            ext = filename.lower().split('.')[-1]
            content = ""

            if ext == 'pdf':
                content = self._extract_pdf(file_path)
                doc_type = "pdf"
            elif ext in ['docx', 'doc']:
                content = self._extract_docx(file_path)
                doc_type = "docx"
            elif ext == 'pptx':
                content = self._extract_pptx(file_path)
                doc_type = "pptx"
            else:
                content = f"[NIEWSPIERANY FORMAT PLIKU: {ext}]"
                doc_type = "unknown"

            sources.append({
                "type": doc_type,
                "filename": filename,
                "content": content
            })
            combined_text_list.append(f"--- Źródło: PLIK {filename} ---\n{content}")

        # 3. Budowanie finalnego słownika
        result = {
            "sources": sources,
            "combined_text": "\n\n".join(combined_text_list)
        }

        # Konwersja do stringa JSON
        json_string = json.dumps(result, ensure_ascii=False, indent=2)

        # 4. Zapisywanie na dysk jako plik
        if save_to_file:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"parsed_input_{timestamp}.json"
            try:
                with open(output_filename, 'w', encoding='utf-8') as f:
                    f.write(json_string)
                print(f"✅ Zapisano plik JSON: {output_filename}")
            except Exception as e:
                print(f"⚠️ Błąd zapisu pliku: {e}")

        # Zwrócenie formatu jako zwykły tekst (zmienna)
        return json_string


# ==========================================
# PRZYKŁAD UŻYCIA
# ==========================================
if __name__ == "__main__":
    parser = BriefInputParser()

    # Dane testowe
    teksty_od_usera = [
        "Cześć, zróbmy kampanię dla nowych butów. Grupa docelowa to młodzież."
    ]
    pliki_od_usera = [
        "TechFest 3.0 - Zadanie.pdf"  # Upewnij się, że ten plik jest w tym samym folderze
    ]

    # Wygenerowanie JSON-a (funkcja zwraca tekst i tworzy plik)
    print("Parsowanie dokumentów wejściowych...")
    wynik_tekstowy = parser.generate_json_payload(
        raw_texts=teksty_od_usera,
        file_paths=pliki_od_usera,
        save_to_file=True
    )

    # Wyświetlenie wynikowego stringa (skrócone dla czytelności w konsoli)
    print("\nWYNIKOWY TEKST JSON (fragment):\n")
    print(wynik_tekstowy[:500] + "\n...\n")