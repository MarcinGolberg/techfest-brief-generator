import os
import uuid
from functools import lru_cache
from typing import Any, Dict, List

from docx import Document
from docx.shared import Pt

from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


def _safe_text(value: Any) -> str:
    if value is None:
        return "-"
    text = str(value).strip()
    return text if text else "-"


def _safe_list(value: Any) -> List[str]:
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    return []


@lru_cache(maxsize=1)
def _get_pdf_font_names() -> tuple[str, str]:
    fonts_dir = os.path.join(os.path.dirname(__file__), "fonts")
    regular_path = os.path.join(fonts_dir, "Graphik-Regular.ttf")
    bold_path = os.path.join(fonts_dir, "Graphik-Bold.ttf")

    if not (os.path.exists(regular_path) and os.path.exists(bold_path)):
        return "Helvetica", "Helvetica-Bold"

    regular_name = "Graphik-Regular-Embedded"
    bold_name = "Graphik-Bold-Embedded"

    registered_fonts = set(pdfmetrics.getRegisteredFontNames())
    if regular_name not in registered_fonts:
        pdfmetrics.registerFont(TTFont(regular_name, regular_path))
    if bold_name not in registered_fonts:
        pdfmetrics.registerFont(TTFont(bold_name, bold_path))

    return regular_name, bold_name


def generate_document_brief_docx(document_brief: Dict[str, Any], output_dir: str = "generated") -> str:
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, f"brief_{uuid.uuid4().hex}.docx")

    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)

    title = _safe_text(document_brief.get("title", "BRIEF NA KAMPANIĘ"))
    doc.add_heading(title, level=0)

    executive_summary = _safe_text(document_brief.get("executive_summary", ""))
    if executive_summary != "-":
        doc.add_heading("EXECUTIVE SUMMARY / STRESZCZENIE", level=2)
        doc.add_paragraph(executive_summary)

    sections = [
        ("BACKGROUND / KONTEKST / OPIS PROJEKTU", document_brief.get("background", "")),
        ("OBJECTIVE / CEL", document_brief.get("objective", "")),
        ("TARGET GROUP (TG) / GRUPA DOCELOWA", document_brief.get("target_group", "")),
        ("ISSUE / PROBLEM", document_brief.get("issue", "")),
        ("INSIGHT / OBSERWACJA", document_brief.get("insight", "")),
        ("CREATIVE CHALLENGE / WYZWANIE KREATYWNE", document_brief.get("creative_challenge", "")),
    ]

    for heading, value in sections:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(_safe_text(value))

    doc.add_heading("MANDATORIES / WYMAGANIA I OGRANICZENIA", level=2)
    mandatories = _safe_list(document_brief.get("mandatories", []))
    if mandatories:
        for item in mandatories:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("-")

    doc.add_heading("Graphic Asset Type", level=2)
    doc.add_paragraph(_safe_text(document_brief.get("graphic_asset_type", "")))

    doc.add_heading("Branding Guidance", level=2)
    branding = _safe_list(document_brief.get("branding_guidance", []))
    if branding:
        for item in branding:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("-")

    doc.add_heading("KEY POINTS", level=2)
    key_points = _safe_list(document_brief.get("key_points", []))
    if key_points:
        for item in key_points:
            doc.add_paragraph(item, style="List Bullet")
    else:
        doc.add_paragraph("-")

    sources_summary = _safe_text(document_brief.get("sources_summary", ""))
    if sources_summary != "-":
        doc.add_heading("ŹRÓDŁA", level=2)
        doc.add_paragraph(sources_summary)

    doc.save(path)
    return path


def generate_document_brief_pdf(document_brief: Dict[str, Any], output_dir: str = "generated") -> str:
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, f"brief_{uuid.uuid4().hex}.pdf")
    regular_font, bold_font = _get_pdf_font_names()

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleCustom",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=20,
        leading=24,
        textColor=HexColor("#4A4AFF"),
        alignment=TA_LEFT,
        spaceAfter=16
    )

    heading_style = ParagraphStyle(
        "HeadingCustom",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=12,
        leading=15,
        textColor=HexColor("#111111"),
        spaceBefore=10,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        "BodyCustom",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=10.5,
        leading=15,
        textColor=HexColor("#222222"),
        spaceAfter=8
    )

    story = []
    story.append(Paragraph(_safe_text(document_brief.get("title", "BRIEF NA KAMPANIĘ")), title_style))

    executive_summary = _safe_text(document_brief.get("executive_summary", ""))
    if executive_summary != "-":
        story.append(Paragraph("EXECUTIVE SUMMARY / STRESZCZENIE", heading_style))
        story.append(Paragraph(executive_summary, body_style))

    sections = [
        ("BACKGROUND / KONTEKST / OPIS PROJEKTU", document_brief.get("background", "")),
        ("OBJECTIVE / CEL", document_brief.get("objective", "")),
        ("TARGET GROUP (TG) / GRUPA DOCELOWA", document_brief.get("target_group", "")),
        ("ISSUE / PROBLEM", document_brief.get("issue", "")),
        ("INSIGHT / OBSERWACJA", document_brief.get("insight", "")),
        ("CREATIVE CHALLENGE / WYZWANIE KREATYWNE", document_brief.get("creative_challenge", "")),
    ]

    for heading, value in sections:
        story.append(Paragraph(heading, heading_style))
        story.append(Paragraph(_safe_text(value), body_style))

    def add_list_section(title: str, items: List[str]):
        story.append(Paragraph(title, heading_style))
        if items:
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(item, body_style)) for item in items],
                    bulletType="bullet",
                    leftIndent=18
                )
            )
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph("-", body_style))

    add_list_section("MANDATORIES / WYMAGANIA I OGRANICZENIA", _safe_list(document_brief.get("mandatories", [])))

    story.append(Paragraph("Graphic Asset Type", heading_style))
    story.append(Paragraph(_safe_text(document_brief.get("graphic_asset_type", "")), body_style))

    add_list_section("Branding Guidance", _safe_list(document_brief.get("branding_guidance", [])))
    add_list_section("KEY POINTS", _safe_list(document_brief.get("key_points", [])))

    sources_summary = _safe_text(document_brief.get("sources_summary", ""))
    if sources_summary != "-":
        story.append(Paragraph("ŹRÓDŁA", heading_style))
        story.append(Paragraph(sources_summary, body_style))

    doc.build(story)
    return path


def generate_brief_file(brief: Dict[str, Any], file_format: str, output_dir: str = "generated") -> str:
    normalized = file_format.strip().lower()

    if normalized == "pdf":
        return generate_document_brief_pdf(brief, output_dir=output_dir)

    if normalized == "docx":
        return generate_document_brief_docx(brief, output_dir=output_dir)

    raise ValueError("Nieobsługiwany format. Użyj 'pdf' albo 'docx'.")
